"""
Generate DigiPermit capstone PowerPoint with live app screenshots.

Usage:
  python scripts/generate-presentation-pptx.py
  python scripts/generate-presentation-pptx.py --skip-screenshots

Output:
  docs/presentation/DigiPermit-Presentation.pptx
  docs/presentation/screenshots/*.png
"""
from __future__ import annotations

import argparse
import asyncio
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs" / "presentation"
SHOT_DIR = OUT_DIR / "screenshots"
PPTX_PATH = OUT_DIR / "DigiPermit-Presentation.pptx"

BASE_URL = "https://mrntuli.github.io/DPermit"
PASSWORD = "Demo@12345"

TEAM = [
    "Mhle (22322987) — System Admin",
    "TV Manqele (22304993) — Employer HR",
    "SK Ngubane (22433510) — Foreign National",
    "S Dube (22322900) — Verification Officer",
    "KS Cebekhulu (22322695) — Immigration Officer",
    "O Luthuli (22325286) — Manager / Auditor",
    "SS Mathonsi (22339617) — University Officer",
    "NN Dlodlo (22325063) — Clinic Administrator",
]

SCREENSHOT_JOBS = [
    ("01-login.png", "/login", None),
    ("02-welcome.png", "/welcome", None),
    ("03-employer-dashboard.png", "/employer/dashboard", ("hr@acmeglobal.demo", PASSWORD)),
    ("04-verify-dashboard.png", "/verification/dashboard", ("verify@digipermit.demo", PASSWORD)),
    ("05-verify-manual.png", "/verification/manual", ("verify@digipermit.demo", PASSWORD)),
    ("06-admin-dashboard.png", "/admin/dashboard", ("admin@digipermit.demo", PASSWORD)),
    ("07-immigration-dashboard.png", "/immigration/dashboard", ("compliance@digipermit.demo", PASSWORD)),
    ("08-manager-dashboard.png", "/manager/dashboard", ("audit@natcompliance.demo", PASSWORD)),
    ("09-fn-permits.png", "/foreign-national/permits", ("james.okonkwo@demo.mail", PASSWORD)),
]

# DigiPermit brand colours
FOREST = (15, 61, 46)
EMERALD = (31, 122, 90)
MINT = (234, 246, 240)
WHITE = (255, 255, 255)
SLATE = (51, 65, 85)
MUTED = (100, 116, 139)


async def capture_screenshots() -> None:
    try:
        from playwright.async_api import async_playwright
    except ImportError:
        import subprocess

        subprocess.check_call([sys.executable, "-m", "pip", "install", "playwright", "-q"])
        subprocess.check_call([sys.executable, "-m", "playwright", "install", "chromium"])
        from playwright.async_api import async_playwright

    SHOT_DIR.mkdir(parents=True, exist_ok=True)

    async def login(page, email: str, password: str) -> None:
        await page.goto(f"{BASE_URL}/login", wait_until="domcontentloaded", timeout=120_000)
        await page.wait_for_timeout(3000)
        email_input = page.locator('input[type="email"]').first
        await email_input.wait_for(state="visible", timeout=60_000)
        await email_input.fill(email)
        await page.locator('input[type="password"]').first.fill(password)
        await page.locator('ion-button[type="submit"]').click()
        await page.wait_for_timeout(6000)

    async with async_playwright() as p:
        browser = await p.chromium.launch()

        for filename, path, creds in SCREENSHOT_JOBS:
            dest = SHOT_DIR / filename
            print(f"  Capturing {filename}...")
            context = await browser.new_context(
                viewport={"width": 1366, "height": 768},
                device_scale_factor=1.25,
            )
            page = await context.new_page()
            try:
                if creds:
                    await login(page, creds[0], creds[1])
                    await page.goto(f"{BASE_URL}{path}", wait_until="domcontentloaded", timeout=120_000)
                else:
                    await page.goto(f"{BASE_URL}{path}", wait_until="domcontentloaded", timeout=120_000)
                await page.wait_for_timeout(5000)
                await page.screenshot(path=str(dest), full_page=False)
            except Exception as exc:
                print(f"    Warning: {filename} failed ({exc})")
            finally:
                await context.close()

        await browser.close()
    print(f"Screenshots saved to {SHOT_DIR}")


def rgb(t):
    from pptx.dml.color import RGBColor
    return RGBColor(*t)


def add_title_slide(prs, title: str, subtitle: str, footer: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(FOREST)

    box = slide.shapes.add_textbox(Inches(0.6), Inches(1.8), Inches(8.8), Inches(1.2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = rgb(WHITE)

    box2 = slide.shapes.add_textbox(Inches(0.6), Inches(3.0), Inches(8.8), Inches(1.5))
    tf2 = box2.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(16)
    p2.font.color.rgb = rgb((200, 230, 215))

    if footer:
        box3 = slide.shapes.add_textbox(Inches(0.6), Inches(5.8), Inches(8.8), Inches(1.2))
        tf3 = box3.text_frame
        p3 = tf3.paragraphs[0]
        p3.text = footer
        p3.font.size = Pt(11)
        p3.font.color.rgb = rgb((180, 210, 195))


def add_section_slide(prs, title: str, bullets: list[str]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # header bar
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.9))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(FOREST)
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.15), Inches(9), Inches(0.7))
    tp = tb.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(26)
    tp.font.bold = True
    tp.font.color.rgb = rgb(WHITE)

    body = slide.shapes.add_textbox(Inches(0.55), Inches(1.15), Inches(9), Inches(5.8))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(15)
        p.font.color.rgb = rgb(SLATE)
        p.space_after = Pt(8)
        p.level = 0


def add_image_slide(prs, title: str, image_path: Path | None, caption: str = "", bullets: list[str] | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.75))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(EMERALD)
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.12), Inches(9), Inches(0.55))
    tp = tb.text_frame.paragraphs[0]
    tp.text = title
    tp.font.size = Pt(22)
    tp.font.bold = True
    tp.font.color.rgb = rgb(WHITE)

    top = 1.0
    if image_path and image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(0.45), Inches(top), width=Inches(9.1))
        top = 5.35
    else:
        ph = slide.shapes.add_shape(1, Inches(0.45), Inches(top), Inches(9.1), Inches(4.2))
        ph.fill.solid()
        ph.fill.fore_color.rgb = rgb(MINT)
        ph.line.color.rgb = rgb(EMERALD)
        pt = slide.shapes.add_textbox(Inches(0.45), Inches(2.5), Inches(9.1), Inches(0.5))
        pt.text_frame.paragraphs[0].text = "[ Screenshot: run generate-presentation-pptx.py ]"
        pt.text_frame.paragraphs[0].font.size = Pt(14)
        pt.text_frame.paragraphs[0].font.color.rgb = rgb(MUTED)
        top = 5.35

    if caption:
        cap = slide.shapes.add_textbox(Inches(0.45), Inches(top), Inches(9.1), Inches(0.4))
        cp = cap.text_frame.paragraphs[0]
        cp.text = caption
        cp.font.size = Pt(10)
        cp.font.italic = True
        cp.font.color.rgb = rgb(MUTED)
        top += 0.35

    if bullets:
        body = slide.shapes.add_textbox(Inches(0.45), Inches(top), Inches(9.1), Inches(1.2))
        tf = body.text_frame
        for i, line in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"• {line}"
            p.font.size = Pt(12)
            p.font.color.rgb = rgb(SLATE)


def add_table_slide(prs, title: str, headers: list[str], rows: list[list[str]]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.75))
    bar.fill.solid()
    bar.fill.fore_color.rgb = rgb(FOREST)
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.12), Inches(9), Inches(0.55))
    tb.text_frame.paragraphs[0].text = title
    tb.text_frame.paragraphs[0].font.size = Pt(22)
    tb.text_frame.paragraphs[0].font.bold = True
    tb.text_frame.paragraphs[0].font.color.rgb = rgb(WHITE)

    cols, row_count = len(headers), len(rows) + 1
    table = slide.shapes.add_table(row_count, cols, Inches(0.4), Inches(1.05), Inches(9.2), Inches(0.35 * row_count)).table

    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(FOREST)
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = rgb(WHITE)

    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(9)
                p.font.color.rgb = rgb(SLATE)


def shot(name: str) -> Path:
    return SHOT_DIR / name


def build_pptx() -> None:
    global Inches, Pt
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        import subprocess
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx", "-q"])
        from pptx import Presentation
        from pptx.util import Inches, Pt

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    team_text = "\n".join(TEAM)
    add_title_slide(
        prs,
        "DigiPermit",
        "Foreign-National Visa & Permit Compliance Monitoring\nDSOF300 / DVSF300 Major Project · DS3 × IS3 × IP2",
        f"June 2026\n{team_text}\nLive: mrntuli.github.io/DPermit",
    )

    add_section_slide(prs, "Problem (As-Is)", [
        "Employers and institutions track permits in spreadsheets and email.",
        "Expiry dates are missed → compliance and legal risk.",
        "No real-time verification or audit trail at checkpoints.",
        "Sensitive data shared without proper access control.",
        "Academic note: DigiPermit monitors compliance — it does not issue official government visas.",
    ])

    add_section_slide(prs, "Solution (To-Be)", [
        "Central database for permits, foreign nationals, and organisations.",
        "Automated daily expiry job → notifications and alerts.",
        "Multi-channel verification: manual, QR scan, RFID/IoT simulation.",
        "Immutable verification_logs for every checkpoint check.",
        "Role-based dashboards, in-app analytics, PDF reports, Power BI (IS3).",
    ])

    add_table_slide(
        prs,
        "Eight user roles",
        ["Role", "Purpose", "Key action"],
        [
            ["System Administrator", "Platform governance", "Orgs, users, logs, analytics"],
            ["Employer HR", "Work-visa compliance", "Register employees, capture permits"],
            ["University Officer", "Study-visa compliance", "Monitor international students"],
            ["Clinic Administrator", "Healthcare records", "Patient permit monitoring"],
            ["Foreign National", "Self-service", "View permit, QR, renewal requests"],
            ["Verification Officer", "Checkpoint", "Scan / verify, audit own scans"],
            ["Immigration Officer", "Simulated review", "Validate pending permits"],
            ["Manager / Auditor", "Oversight", "Read-only analytics & insights"],
        ],
    )

    add_section_slide(prs, "System architecture (IP2 + IS3)", [
        "Presentation: Ionic Angular (TypeScript) — mobile-first government UI.",
        "Business logic: Node.js + Express REST API (40+ endpoints).",
        "Data: Supabase PostgreSQL + Supabase Auth (JWT).",
        "Security: RBAC on routes, API middleware, row-level organisation scope.",
        "Analytics: Chart.js dashboards, SQL views, Power BI connection.",
    ])

    add_section_slide(prs, "IS3 design artefacts", [
        "BPMN 2.0 As-Is and To-Be process models (manual → automated expiry).",
        "3NF ERD — 12 tables, foreign keys, analytics views.",
        "IoT connectivity map — Packet Tracer + API simulation.",
        "AI workflow — rule-based suspicious activity + executive insights API.",
    ])

    add_image_slide(
        prs,
        "Login & secure access",
        shot("01-login.png"),
        "JWT authentication via Supabase Auth — each user sees only authorised menus.",
        ["Demo password for all users: Demo@12345"],
    )

    add_image_slide(
        prs,
        "Employer HR — compliance dashboard",
        shot("03-employer-dashboard.png"),
        "Organisation-scoped view: employees, permits, expiring visas, alerts.",
        ["Registers foreign employees and captures work permits (UC-01)."],
    )

    add_image_slide(
        prs,
        "Foreign national — digital permit",
        shot("09-fn-permits.png"),
        "Self-service: view own permits, QR code, notifications, renewal requests.",
        ["Employee shows QR on phone; officer scans at checkpoint."],
    )

    add_image_slide(
        prs,
        "Verification officer — checkpoint dashboard",
        shot("04-verify-dashboard.png"),
        "Live analytics filtered to officer's own scans; PDF export of current view.",
        ["Separation of duties: admin governs platform; officers perform scans."],
    )

    add_image_slide(
        prs,
        "Manual & QR verification",
        shot("05-verify-manual.png"),
        "Ten result states: Valid, Expiring soon, Expired, Revoked, Not found, etc.",
        ["Try WP-2024-ACME-001 (valid) and WP-2023-ACME-003 (expired) in live demo."],
    )

    add_image_slide(
        prs,
        "System administrator — platform overview",
        shot("06-admin-dashboard.png"),
        "Cross-organisation statistics, verification logs, alerts, user management.",
        ["Admin reviews audit trail — does not replace checkpoint verification role."],
    )

    add_image_slide(
        prs,
        "Immigration officer — simulated review",
        shot("07-immigration-dashboard.png"),
        "Validates pending permits; reviews suspicious activity and renewals.",
        ["Academic simulation — not Department of Home Affairs."],
    )

    add_image_slide(
        prs,
        "Manager / auditor — executive view",
        shot("08-manager-dashboard.png"),
        "Read-only analytics and AI insight summary for compliance reporting.",
    )

    add_section_slide(prs, "IoT & RFID simulation (IS3)", [
        "Devices registered in platform (e.g. IOT-ACME-GATE-01).",
        "RFID scan POSTs to /api/iot/simulate → verification_logs + iot_scan_events.",
        "Cisco Packet Tracer topology documents network path (see lab guide).",
        "Demonstrates live data feeding 3NF database — not offline spreadsheet.",
    ])

    add_section_slide(prs, "Security & auditability", [
        "Role-based access control (8 roles) on frontend and API.",
        "Verification logs are immutable; tied to officer and timestamp.",
        "Passport masking on verification responses.",
        "Organisation-scoped data for HR, university, clinic roles.",
    ])

    add_section_slide(prs, "Power BI & reporting (IS3)", [
        "SQL views: vw_expiring_permits, verification trends, permit status.",
        "Power BI Desktop connects to Supabase (live or CSV export).",
        "In-app Download PDF report — role-specific, live data at export time.",
        "See docs/power-bi/POWER-BI-LIVE-SUPABASE.md for setup steps.",
    ])

    add_section_slide(prs, "Agile delivery (DS3)", [
        "Scrum over 4 sprints / 6 weeks.",
        "Sprint 1: Database + API foundation.",
        "Sprint 2: Verification, IoT, employer UI.",
        "Sprint 3: All roles, analytics, AI insights.",
        "Sprint 4: Testing, documentation, deployment, presentation.",
        "GitHub: github.com/MrNtuli/DPermit",
    ])

    add_table_slide(
        prs,
        "Live demo credentials",
        ["Role", "Email", "Password"],
        [
            ["System Admin", "admin@digipermit.demo", "Demo@12345"],
            ["Employer HR", "hr@acmeglobal.demo", "Demo@12345"],
            ["Foreign National", "james.okonkwo@demo.mail", "Demo@12345"],
            ["Verification Officer", "verify@digipermit.demo", "Demo@12345"],
            ["Immigration Officer", "compliance@digipermit.demo", "Demo@12345"],
            ["Manager", "audit@natcompliance.demo", "Demo@12345"],
        ],
    )

    add_table_slide(
        prs,
        "Speaker allocation (suggested)",
        ["Step", "Speaker", "Content"],
        [
            ["1–2", "TV Manqele", "Problem, HR capture"],
            ["3", "SK Ngubane + S Dube", "QR show + scan (rehearse!)"],
            ["4", "KS Cebekhulu", "Immigration validate"],
            ["5", "Mhle", "Admin logs & analytics"],
            ["6", "SS Mathonsi / NN Dlodlo", "University & clinic"],
            ["7", "O Luthuli", "Manager insights"],
            ["8", "Any member", "Limitation statement + Q&A"],
        ],
    )

    add_title_slide(prs, "Thank you", "Questions?", "DigiPermit — Government Compliance Portal (Academic Prototype)")

    prs.save(str(PPTX_PATH))
    print(f"OK  {PPTX_PATH} ({PPTX_PATH.stat().st_size / 1024:.0f} KB)")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--skip-screenshots", action="store_true")
    args = parser.parse_args()

    if not args.skip_screenshots:
        print("Capturing live app screenshots...")
        asyncio.run(capture_screenshots())
    else:
        print("Skipping screenshots.")

    print("Building PowerPoint...")
    build_pptx()


if __name__ == "__main__":
    main()
